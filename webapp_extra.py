"""Mini app qo'shimcha imkoniyatlari: dori katalogi, «kelganda xabar berish», hisob-kitob yaratuvchi,
pauza, operator kunlik maqsadi va adminlarga kunlik hisobot.

webapp.build_app() ichida register(app) chaqiriladi; fon vazifalari start_background() orqali.
"""
import io
import base64
import asyncio
import logging
import html as _htm
from datetime import datetime, timedelta

from aiohttp import web

from config import now_local
from database import queries as q

logger = logging.getLogger("bot")


def _w():
    import webapp
    return webapp


def _json(data, status=200):
    return web.json_response(data, status=status)


async def _body(request):
    try:
        return await request.json()
    except Exception:
        return {}


def _money(n) -> str:
    try:
        return f"{int(round(float(n))):,}".replace(",", " ")
    except (TypeError, ValueError):
        return "0"


def _prod_json(r, waits=None, bnames=None):
    bid = (r["branch_id"] if "branch_id" in r.keys() else None) or 0
    return {"id": r["id"], "name": r["name"], "price": r["price"] or 0, "unit": r["unit"] or "",
            "note": r["note"] or "", "in_stock": bool(r["in_stock"]),
            "branch_id": bid, "branch": (bnames or {}).get(bid, "") if bid else "",
            "waits": (waits or {}).get(r["id"], 0)}


async def _branch_names() -> dict:
    return {b["id"]: b["name"] for b in await q.list_branches()}


async def _ai_enabled() -> bool:
    import ai_catalog
    return await ai_catalog.enabled()


def _br_arg(v):
    """'all'/'' -> None (hammasi); '0' -> 0 (umumiy); '5' -> 5"""
    v = str(v if v is not None else "").strip()
    if v in ("", "all"):
        return None
    try:
        return max(0, int(v))
    except ValueError:
        return None


# ================================================================
#                        DORI KATALOGI
# ================================================================
async def api_products(request):
    """Operator: chat ichida dori qidirish (nomi, narxi, mavjudligi)."""
    op, _ = await _w()._auth_op(request, request.query)
    if not op:
        return _json({"ok": False}, 401)
    term = (request.query.get("q") or "").strip()
    # Chat ochiq bo'lsa — mijoz filiali katalogi + umumiy katalog
    branch, bname = None, ""
    try:
        oid = int(request.query.get("order_id") or 0)
    except ValueError:
        oid = 0
    if oid:
        shared = bool(request.query.get("shared"))
        order, denied = await _w()._operator_order(op, oid, allow_unassigned_new=not shared,
                                                   allow_shared=shared)
        if denied is not None:
            return denied
        user = await q.get_user(order["user_id"])
        bid = order["branch_id"] or (user["branch_id"] if user else None)
        if bid and await q.products_count(bid):
            branch = bid
    bnames = await _branch_names()
    if branch:
        bname = bnames.get(branch, "")
    rows = await q.search_products(term, 40, branch=branch, with_common=True)
    waits = await q.stock_wait_counts()
    return _json({"ok": True, "items": [_prod_json(r, waits, bnames) for r in rows],
                  "branch": bname, "total": await q.products_count(branch, with_common=True)})


async def api_stock_wait(request):
    """Dori yo'q — kelganda mijozga avtomatik xabar boradi."""
    body = await _body(request)
    op, _ = await _w()._auth_op(request, body)
    if not op:
        return _json({"ok": False}, 401)
    try:
        order_id = int(body.get("order_id"))
        pid = int(body.get("product_id"))
    except (TypeError, ValueError):
        return _json({"ok": False}, 400)
    order, denied = await _w()._operator_order(op, order_id,
                                               allow_shared=bool(body.get("shared")))
    if denied is not None:
        return denied
    prod = await q.get_product(pid)
    if not order or not prod:
        return _json({"ok": False, "error": "topilmadi"}, 404)
    if prod["in_stock"]:
        return _json({"ok": False, "error": "Bu dori hozir mavjud — navbatga qo'shilmadi"})
    user = await q.get_user(order["user_id"])
    order_branch = order["branch_id"] or (user["branch_id"] if user else None)
    product_branch = prod["branch_id"] if "branch_id" in prod.keys() else None
    if product_branch and product_branch != order_branch:
        return _json({"ok": False, "error": "Dori mijoz tanlagan filial katalogiga tegishli emas"}, 409)
    added = await q.add_stock_wait(pid, order["user_id"], order_id, op["id"])
    from utils import cbot
    client = cbot()
    if client and body.get("tell_client", True):
        lang = await q.get_lang(order["user_id"])
        txt = (f"🔔 «<b>{_htm.escape(prod['name'])}</b>» hozircha yo'q. Dorixonamizga kelishi bilan sizga shu "
               f"yerda avtomatik xabar beramiz.") if lang != "ru" else (
               f"🔔 «<b>{_htm.escape(prod['name'])}</b>» сейчас нет в наличии. Как только поступит — "
               f"мы сразу сообщим вам здесь.")
        try:
            snt = await client.send_message(order["user_id"], txt)
            await q.add_message(order_id, "operator", "text", txt, None, None, client_msg_id=snt.message_id)
        except Exception:
            pass
    await q.audit(f"op:{op['name']}", "stock_wait", order_id, {"product": prod["name"]})
    return _json({"ok": True, "queued": added,
                  "info": "Dori kelganda mijozga xabar boradi" if added
                          else "Mijoz bu dori uchun avval navbatga qo'yilgan"})


async def notify_stock_arrived(pid):
    """Dori «bor» bo'lganda — kutayotgan mijozlarga xabar."""
    prod = await q.get_product(pid)
    if not prod or not prod["in_stock"]:
        return 0
    rows = await q.stock_waits_for_product(pid)
    if not rows:
        return 0
    from utils import cbot
    client = cbot()
    if not client:
        return 0
    sent = 0
    nm = _htm.escape(prod["name"])
    for w in rows:
        lang = await q.get_lang(w["user_id"])
        if lang == "ru":
            price = f"\n💰 Цена: {_money(prod['price'])} сум" if prod["price"] else ""
            txt = f"✅ Хорошая новость! «<b>{nm}</b>» поступил в аптеку.{price}\n\nНапишите сюда, чтобы оформить заказ 💊"
        else:
            price = f"\n💰 Narxi: {_money(prod['price'])} so'm" if prod["price"] else ""
            txt = f"✅ Xushxabar! Siz so'ragan «<b>{nm}</b>» dorixonamizga keldi.{price}\n\nBuyurtma berish uchun shu yerga yozing 💊"
        try:
            sent_msg = await client.send_message(w["user_id"], txt)
            sent += 1
            await q.delete_stock_wait(w["id"])
            order = await q.get_order(w["order_id"]) if w["order_id"] else None
            if order and order["status"] in ("new", "in_progress"):
                await q.add_message(order["id"], "operator", "text", txt, None, None,
                                    client_msg_id=sent_msg.message_id)
        except Exception:
            pass
    await q.audit("system", "stock_arrived", None, {"product": prod["name"], "notified": sent})
    return sent


# ---------------- Admin: katalog boshqaruvi ----------------
async def api_admin_products(request):
    if not await _w()._auth_admin(request, request.query):
        return _json({"ok": False}, 401)
    term = (request.query.get("q") or "").strip()
    try:
        page = int(request.query.get("page") or 0)
    except ValueError:
        page = 0
    br = _br_arg(request.query.get("br"))
    rows = await q.search_products(term, 100000, branch=br)
    waits = await q.stock_wait_counts()
    total = len(rows)
    rows = rows[page * 50:(page + 1) * 50]
    bnames = await _branch_names()
    counts = await q.product_branch_counts()
    return _json({"ok": True, "total": total, "page": page, "br": "all" if br is None else br,
                  "items": [_prod_json(r, waits, bnames) for r in rows],
                  "waiting": sum(waits.values()),
                  "all_cnt": sum(counts.values()), "common_cnt": counts.get(0, 0),
                  "ai": await _ai_enabled(),
                  "branches": [{"id": i, "name": n, "cnt": counts.get(i, 0)} for i, n in bnames.items()]})


async def api_admin_product_save(request):
    body = await _body(request)
    tg = await _w()._auth_admin(request, body)
    if not tg:
        return _json({"ok": False}, 401)
    name = str(body.get("name", "")).strip()[:200]
    if not name:
        return _json({"ok": False, "error": "Nomi bo'sh"})
    try:
        price = int(float(str(body.get("price") or 0).replace(" ", "").replace(",", ".")))
    except ValueError:
        return _json({"ok": False, "error": "Narx noto'g'ri"})
    pid = int(body.get("id") or 0) or None
    old = await q.get_product(pid) if pid else None
    stock = 1 if body.get("in_stock", True) else 0
    pid = await q.save_product(pid, name, price, str(body.get("unit") or "")[:30],
                               str(body.get("note") or "")[:300], stock, _br_arg(body.get("br")) or None)
    n = 0
    if stock and old is not None and not old["in_stock"]:
        n = await notify_stock_arrived(pid)
    return _json({"ok": True, "id": pid, "notified": n})


async def api_admin_product_stock(request):
    body = await _body(request)
    tg = await _w()._auth_admin(request, body)
    if not tg:
        return _json({"ok": False}, 401)
    try:
        pid = int(body.get("id"))
    except (TypeError, ValueError):
        return _json({"ok": False}, 400)
    stock = bool(body.get("in_stock"))
    await q.set_product_stock(pid, stock)
    n = await notify_stock_arrived(pid) if stock else 0
    return _json({"ok": True, "notified": n})


async def api_admin_product_del(request):
    body = await _body(request)
    if not await _w()._auth_admin(request, body):
        return _json({"ok": False}, 401)
    await q.delete_product(int(body.get("id")))
    return _json({"ok": True})


_HDR = {
    "name": ("nomi", "nom", "name", "dori", "наименование", "название", "товар", "препарат"),
    "price": ("narx", "narxi", "price", "цена", "стоимость", "summa", "сумма"),
    "unit": ("birlik", "unit", "ед", "единица", "upakovka", "упаковка"),
    "note": ("izoh", "note", "tavsif", "описание", "производитель", "ishlab"),
    "stock": ("mavjud", "stock", "qoldiq", "остаток", "наличие", "bor"),
}


def _parse_xlsx(raw: bytes):
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    ws = wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    # sarlavha qatorini topamiz (birinchi 10 qatordan)
    cols, start = {}, 0
    for i, row in enumerate(rows[:10]):
        found = {}
        for j, cell in enumerate(row):
            v = str(cell or "").strip().lower()
            for k, keys in _HDR.items():
                if k not in found and any(v.startswith(x) for x in keys):
                    found[k] = j
        if "name" in found:
            cols, start = found, i + 1
            break
    if not cols:
        cols, start = {"name": 0, "price": 1}, 0
    out = []
    for row in rows[start:]:
        def g(k):
            j = cols.get(k)
            return row[j] if j is not None and j < len(row) else None
        name = str(g("name") or "").strip()
        if not name:
            continue
        try:
            price = int(float(str(g("price") or 0).replace(" ", "").replace(",", ".") or 0))
        except ValueError:
            price = 0
        st = g("stock")
        if st is None or str(st).strip() == "":
            stock = 1
        else:
            sv = str(st).strip().lower()
            try:
                stock = 1 if float(sv.replace(",", ".")) > 0 else 0
            except ValueError:
                stock = 0 if sv in ("yo'q", "yoq", "нет", "no", "0", "-", "false") else 1
        out.append((name[:200], price, str(g("unit") or "")[:30], str(g("note") or "")[:300], stock))
    return out


async def api_admin_products_import(request):
    body = await _body(request)
    tg = await _w()._auth_admin(request, body)
    if not tg:
        return _json({"ok": False}, 401)
    try:
        raw = base64.b64decode(str(body.get("data") or "").split(",")[-1])
        items = _parse_xlsx(raw)
    except Exception:
        logger.exception("products import")
        return _json({"ok": False, "error": "Faylni o'qib bo'lmadi. .xlsx formatida yuklang."})
    if not items:
        return _json({"ok": False, "error": "Faylda dori topilmadi (1-ustun: Nomi, 2-ustun: Narxi)"})
    br = _br_arg(body.get("br")) or None
    if br and not await q.get_branch(br):
        return _json({"ok": False, "error": "Filial topilmadi"})
    added, updated = await q.import_products(items, replace=bool(body.get("replace")), branch_id=br)
    # kelgan dorilar bo'yicha kutayotganlarga xabar
    notified = 0
    waits = await q.stock_wait_counts()
    for pid in waits:
        p = await q.get_product(pid)
        if p and p["in_stock"]:
            notified += await notify_stock_arrived(pid)
    await q.audit(f"admin:{tg}", "products_import", None, {"added": added, "updated": updated, "branch": br or 0})
    return _json({"ok": True, "added": added, "updated": updated, "notified": notified})


_SAMPLE = [("Paratsetamol 500mg", 12000, "quti", "10 tabletka", "bor"),
           ("Aspirin Kardio 100mg", 38000, "quti", "Bayer", "bor"),
           ("Vitamin D3 2000 IU", 95000, "quti", "60 kapsula", "yo'q"),
           ("Nurofen sirop 100ml", 54000, "dona", "bolalar uchun", "bor")]


def build_catalog_xlsx(rows, branch_name="") -> bytes:
    """Katalog Excel fayli: namuna (rows=None) yoki joriy katalog. Import xuddi shu formatni o'qiydi."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.worksheet.datavalidation import DataValidation
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Katalog"
    ws.append([f"Gulnora Farm — {branch_name or 'Umumiy katalog (barcha filiallar)'}"])
    ws.append(["Faqat 4-qatordan pastini to'ldiring. Sarlavhalarni o'zgartirmang."])
    ws.append([])
    ws.append(["Nomi", "Narxi", "Birlik", "Izoh", "Mavjud"])
    ws["A1"].font = Font(bold=True, size=14, color="1B7F4B")
    ws["A2"].font = Font(italic=True, size=10, color="888888")
    head_fill = PatternFill("solid", fgColor="1B7F4B")
    thin = Side(style="thin", color="D0D7DE")
    for c in ws[4]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = head_fill
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = Border(bottom=thin)
    ws.row_dimensions[4].height = 22
    data = rows if rows is not None else _SAMPLE
    for name, price, unit, note, stock in data:
        if not isinstance(stock, str):
            stock = "bor" if stock else "yo'q"
        ws.append([name, price or 0, unit or "", note or "", stock])
    last = max(5, ws.max_row) + 500
    for r in ws.iter_rows(min_row=5, max_row=ws.max_row, min_col=2, max_col=2):
        for c in r:
            c.number_format = "# ##0"
    dv = DataValidation(type="list", formula1='"bor,yo\'q"', allow_blank=True,
                        error="«bor» yoki «yo'q» tanlang", errorTitle="Mavjud")
    ws.add_data_validation(dv)
    dv.add(f"E5:E{last}")
    for col, wdt in zip("ABCDE", (38, 13, 11, 30, 10)):
        ws.column_dimensions[col].width = wdt
    ws.freeze_panes = "A5"
    # Yo'riqnoma
    h = wb.create_sheet("Yo'riqnoma")
    for line in ["Qanday to'ldiriladi:",
                 "• Nomi — dori nomi (majburiy). Bir xil nomli dori qayta yuklansa, yangilanadi.",
                 "• Narxi — so'mda, faqat raqam (masalan 12000).",
                 "• Birlik — quti, dona, ml ... (ixtiyoriy).",
                 "• Izoh — ishlab chiqaruvchi, dozasi (ixtiyoriy).",
                 "• Mavjud — «bor» yoki «yo'q» (1/0 ham bo'ladi). Bo'sh qolsa — «bor».",
                 "",
                 "Yuklash: Admin panel → Sozlamalar → Dori katalogi → filialni tanlang → «Excel yuklash».",
                 "«Mavjud» bo'lib qolgan dorini kutayotgan mijozlarga avtomatik xabar boradi."]:
        h.append([line])
    h["A1"].font = Font(bold=True, size=13)
    h.column_dimensions["A"].width = 100
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _xlsx_resp(raw, fname):
    return web.Response(body=raw, headers={
        "Content-Type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "Content-Disposition": f'attachment; filename="{fname}"'})


async def api_admin_products_template(request):
    """Namuna Excel fayl (havola orqali yuklab olish — avtorizatsiyasiz, faqat namuna)."""
    return _xlsx_resp(build_catalog_xlsx(None), "katalog_namuna.xlsx")


async def api_admin_products_file(request):
    """Namuna yoki tanlangan filialning joriy katalogi — adminning Telegram chatiga fayl bo'lib boradi
    (telefonda ham ishonchli yuklab olinadi, keyin to'ldirib qayta yuklanadi)."""
    body = await _body(request)
    tg = await _w()._auth_admin(request, body)
    if not tg:
        return _json({"ok": False}, 401)
    br = _br_arg(body.get("br"))
    kind = body.get("kind") or "template"
    bnames = await _branch_names()
    bname = bnames.get(br, "") if br else ""
    slug = "".join(ch if ch.isalnum() else "_" for ch in (bname or "umumiy")).strip("_")[:30] or "katalog"
    if kind == "export":
        rows = await q.search_products("", 100000, branch=br if br is not None else 0)
        data = [(r["name"], r["price"], r["unit"], r["note"], r["in_stock"]) for r in rows]
        raw = build_catalog_xlsx(data, bname)
        fname = f"katalog_{slug}.xlsx"
        cap = f"📦 Joriy katalog: <b>{_htm.escape(bname or 'Umumiy')}</b> — {len(data)} ta dori.\n" \
              f"Tahrirlab, admin panelda shu filialni tanlab qayta yuklang."
    else:
        raw = build_catalog_xlsx(None, bname)
        fname = f"katalog_namuna_{slug}.xlsx"
        cap = "📄 Dori katalogi namunasi. To'ldirib, admin panel → Dori katalogi → «Excel yuklash» orqali yuklang."
    from aiogram.types import BufferedInputFile
    from utils import cbot
    bot = cbot()
    if not bot:
        return _json({"ok": False, "error": "bot tayyor emas"})
    try:
        await bot.send_document(tg, BufferedInputFile(raw, fname), caption=cap)
    except Exception:
        logger.exception("products file -> admin")
        return _json({"ok": False, "error": "Telegramga yuborib bo'lmadi — botni ishga tushiring (/start)"})
    return _json({"ok": True, "info": "Fayl Telegram chatingizga yuborildi ✅"})


# ================================================================
#            AI (Gemini / Groq): Excel'ni tahlil qilib filiallarga taqsimlash
# ================================================================
async def api_admin_ai_settings(request):
    if not await _w()._auth_admin(request, request.query):
        return _json({"ok": False}, 401)
    import ai_catalog as ai
    out = {}
    for p in ai.PROVIDERS:
        key = (await q.get_setting(f"{p}_api_key", "") or "").strip()
        out[p] = {"has_key": bool(key), "key_mask": ai.mask_key(key),
                  "model": (await q.get_setting(f"{p}_model", "") or "") if key else ""}
    return _json({"ok": True, **out})


async def api_admin_ai_settings_save(request):
    """{provider: gemini|groq, key} — kalit tekshiriladi, model avtomatik tanlanadi. {provider, clear} — o'chirish."""
    body = await _body(request)
    tg = await _w()._auth_admin(request, body)
    if not tg:
        return _json({"ok": False}, 401)
    import ai_catalog as ai
    p = str(body.get("provider") or "")
    if p not in ai.PROVIDERS:
        return _json({"ok": False, "error": "provider"}, 400)
    if body.get("clear"):
        await q.set_setting(f"{p}_api_key", "")
        await q.set_setting(f"{p}_model", "")
        await q.audit(f"admin:{tg}", "settings", None, {f"{p}_api_key": "o'chirildi"})
        return _json({"ok": True, "info": f"{ai.LABEL[p]} kaliti o'chirildi"})
    key = str(body.get("key") or "").strip()
    if len(key) < 20 or any(ch.isspace() for ch in key):
        return _json({"ok": False, "error": "API kalit noto'g'ri ko'rinishda"})
    try:
        model = await ai.check_key(p, key)        # kalit ishlaydimi + eng mos model
    except ai.AIError as e:
        return _json({"ok": False, "error": str(e)})
    await q.set_setting(f"{p}_api_key", key)
    await q.set_setting(f"{p}_model", model)
    await q.audit(f"admin:{tg}", "settings", None, {f"{p}_api_key": ai.mask_key(key), f"{p}_model": model})
    return _json({"ok": True, "model": model, "info": f"{ai.LABEL[p]} ulandi ✅\nModel: {model}"})


async def api_admin_ai_test(request):
    body = await _body(request)
    if not await _w()._auth_admin(request, body):
        return _json({"ok": False}, 401)
    import ai_catalog as ai
    p = str(body.get("provider") or "")
    key = (await q.get_setting(f"{p}_api_key", "") or "").strip() if p in ai.PROVIDERS else ""
    if not key:
        return _json({"ok": False, "error": "API kalit kiritilmagan"})
    try:
        model = await ai.check_key(p, key)
    except ai.AIError as e:
        return _json({"ok": False, "error": str(e)})
    await q.set_setting(f"{p}_model", model)
    return _json({"ok": True, "info": f"{ai.LABEL[p]} ishlayapti ✅\nModel: {model}"})


async def api_admin_products_ai(request):
    """1-qadam: Excel -> Gemini tahlili -> reja (hali hech narsa saqlanmaydi, admin ko'rib tasdiqlaydi)."""
    body = await _body(request)
    tg = await _w()._auth_admin(request, body)
    if not tg:
        return _json({"ok": False}, 401)
    import ai_catalog as ai
    try:
        raw = base64.b64decode(str(body.get("data") or "").split(",")[-1])
    except Exception:
        return _json({"ok": False, "error": "Fayl o'qilmadi"})
    if len(raw) > 8 * 1024 * 1024:
        return _json({"ok": False, "error": "Fayl juda katta (maks 8 MB)"})
    try:
        res = await ai.analyze_catalog_excel(raw)
    except ai.AIError as e:
        return _json({"ok": False, "error": str(e)})
    except Exception:
        logger.exception("products ai")
        return _json({"ok": False, "error": "AI tahlili bajarilmadi"})
    if not res["items"]:
        return _json({"ok": False, "error": "AI faylda dori topmadi"})
    bnames = await _branch_names()
    cnt = {}
    for it in res["items"]:
        cnt[it["branch_id"]] = cnt.get(it["branch_id"], 0) + 1
    groups = [{"id": 0, "name": "Umumiy (barcha filiallar)", "cnt": cnt.get(0, 0)}] + \
             [{"id": i, "name": n, "cnt": cnt.get(i, 0)} for i, n in bnames.items()] + \
             [{"id": -1, "name": "Filiali aniqlanmadi", "cnt": cnt.get(-1, 0)}]
    await q.audit(f"admin:{tg}", "products_ai", None, {"items": len(res["items"])})
    return _json({"ok": True, "items": res["items"], "groups": [g for g in groups if g["cnt"]],
                  "truncated": res["truncated"], "failed_chunks": res["failed_chunks"]})


async def api_admin_products_ai_apply(request):
    """2-qadam: tasdiqlangan rejani filiallar bo'yicha katalogga yozish."""
    body = await _body(request)
    tg = await _w()._auth_admin(request, body)
    if not tg:
        return _json({"ok": False}, 401)
    items = body.get("items") or []
    if not isinstance(items, list) or not items:
        return _json({"ok": False, "error": "Ro'yxat bo'sh"})
    valid = set((await _branch_names()).keys())
    by_br = {}
    for it in items[:20000]:
        try:
            bid = int(it.get("branch_id") or 0)
            price = max(0, int(float(it.get("price") or 0)))
        except (TypeError, ValueError, AttributeError):
            continue
        name = str(it.get("name") or "").strip()[:200]
        if not name or bid < 0 or (bid and bid not in valid):
            continue      # filiali aniqlanmaganlar qo'shilmaydi
        by_br.setdefault(bid, []).append((name, price, str(it.get("unit") or "")[:30],
                                          str(it.get("note") or "")[:300], 1 if it.get("in_stock", True) else 0))
    if not by_br:
        return _json({"ok": False, "error": "Qo'shiladigan dori yo'q (filiali aniqlanmaganlarini filialga biriktiring)"})
    added = updated = 0
    for bid, rows in by_br.items():
        a, u = await q.import_products(rows, replace=False, branch_id=bid or None)
        added += a
        updated += u
    notified = 0
    for pid in await q.stock_wait_counts():
        p = await q.get_product(pid)
        if p and p["in_stock"]:
            notified += await notify_stock_arrived(pid)
    await q.audit(f"admin:{tg}", "products_ai_apply", None,
                  {"added": added, "updated": updated, "branches": len(by_br)})
    return _json({"ok": True, "added": added, "updated": updated, "notified": notified, "branches": len(by_br)})


# ================================================================
#                        HISOB-KITOB YARATUVCHI
# ================================================================
def build_bill_text(items, delivery=0, note=""):
    """[{name, qty, price}] -> chiroyli hisob-kitob matni (oddiy matn) va jami summa."""
    lines, total = [], 0
    for i, it in enumerate(items, 1):
        name = str(it.get("name") or "").strip()[:120]
        if not name:
            continue
        try:
            qty = max(1, int(float(it.get("qty") or 1)))
            price = max(0, int(float(it.get("price") or 0)))
        except (TypeError, ValueError):
            qty, price = 1, 0
        s = qty * price
        total += s
        lines.append(f"{i}. {name} — {qty} × {_money(price)} = {_money(s)} so'm")
    try:
        delivery = max(0, int(float(delivery or 0)))
    except (TypeError, ValueError):
        delivery = 0
    if delivery:
        lines.append(f"🚚 Yetkazib berish: {_money(delivery)} so'm")
        total += delivery
    lines.append("────────────")
    lines.append(f"💰 Jami: {_money(total)} so'm")
    if note:
        lines.append(f"📝 {str(note).strip()[:300]}")
    return "\n".join(lines), total


# ================================================================
#                              PAUZA
# ================================================================
_PAUSE_REASONS = ["Mijoz keyinroq yozadi", "Mijoz ertaga yozadi", "Dori kutilmoqda",
                  "Mijoz o'ylab ko'radi", "To'lov kutilmoqda", "Narx aniqlanmoqda"]


def _fmt_until(s):
    if not s:
        return ""
    try:
        d = datetime.strptime(s[:16], "%Y-%m-%d %H:%M")
        n = now_local().replace(tzinfo=None)
        if d.date() == n.date():
            return "bugun " + d.strftime("%H:%M")
        if d.date() == (n + timedelta(days=1)).date():
            return "ertaga " + d.strftime("%H:%M")
        return d.strftime("%d.%m %H:%M")
    except ValueError:
        return s[:16]


async def api_pause(request):
    body = await _body(request)
    op, _ = await _w()._auth_op(request, body)
    if not op:
        return _json({"ok": False}, 401)
    try:
        order_id = int(body.get("order_id"))
    except (TypeError, ValueError):
        return _json({"ok": False}, 400)
    order, denied = await _w()._operator_order(op, order_id,
                                               allow_shared=bool(body.get("shared")))
    if denied is not None:
        return denied
    if order["status"] not in ("new", "in_progress"):
        return _json({"ok": False, "error": "Murojaat yopilgan"})
    if body.get("resume"):
        mins = await q.resume_order(order_id, "operator")
        await q.audit(f"op:{op['name']}", "pause_end", order_id, {"minutes": round(mins)})
        return _json({"ok": True, "info": "Suhbat davom ettirildi"})
    reason = str(body.get("reason") or "").strip()[:120] or "Pauza"
    until = None
    try:
        mins = int(body.get("minutes") or 0)
    except (TypeError, ValueError):
        mins = 0
    if mins > 0:
        until = (now_local() + timedelta(minutes=min(mins, 14 * 1440))).strftime("%Y-%m-%d %H:%M:%S")
    try:
        remind = max(0, min(int(body.get("remind") or 0), 7 * 1440))
    except (TypeError, ValueError):
        remind = 0
    await q.pause_order(order_id, op["id"], reason, until, remind or None)
    if until:
        # vaqt kelganda operatorga botda eslatma (eslatmalar tizimi orqali, faqat ish vaqtida)
        await q.add_reminder(op["id"], order_id, until, f"⏸ Pauza tugadi: {reason}")
    # Pauza faqat ichki ish holati: mijozga xabar yuborilmaydi.
    await q.audit(f"op:{op['name']}", "pause", order_id,
                  {"reason": reason, "until": until, "client_notified": False})
    return _json({"ok": True, "info": "Suhbat pauzaga qo'yildi — bu vaqt statistikaga kirmaydi",
                  "until": until})


def pause_json(order, pauses=None):
    keys = order.keys()
    pa = order["paused_at"] if "paused_at" in keys else None
    res = {"paused": bool(pa), "paused_at": pa or "",
           "remind_min": (order["pause_remind_min"] or 0) if "pause_remind_min" in keys else 0, "until": (order["paused_until"] or "") if "paused_until" in keys else "",
           "until_label": _fmt_until(order["paused_until"]) if "paused_until" in keys else "",
           "total_min": round(order["paused_total_min"] or 0) if "paused_total_min" in keys else 0}
    if pauses is not None:
        cur = [p for p in pauses if not p["ended_at"]]
        res["reason"] = cur[-1]["reason"] if cur else ""
        res["history"] = [{"reason": p["reason"] or "", "start": p["started_at"], "end": p["ended_at"] or "",
                           "minutes": round(p["minutes"] or 0), "by": p["ended_by"] or ""} for p in pauses]
    return res


def _rem_label(m):
    m = int(m or 0)
    if m % 1440 == 0:
        return f"{m // 1440} kun"
    if m % 60 == 0:
        return f"{m // 60} soat"
    return f"{m} daqiqa"


async def _send_pause_reminder(r):
    """Pauzadagi suhbat haqida operatorga eslatma — faqat uning ish vaqtida."""
    op = await q.get_operator(r["operator_id"]) if r["operator_id"] else None
    if not op or not op["telegram_id"]:
        await q.set_pause_next_remind(r["id"], r["pause_remind_min"])
        return
    from utils import operator_in_hours, cbot
    if not operator_in_hours(op)[0]:
        return            # ish vaqtidan tashqarida — yubormaymiz; ish vaqti boshlanishi bilan yuboriladi
    import botreg
    b = (botreg.get_operator_bot(op["bot_id"]) if op["bot_id"] else None) or cbot()
    since = _fmt_until(r["paused_at"])
    text = (f"⏸ <b>Eslatma: suhbat pauzada</b> — #{r['id']} ({_htm.escape(r['full_name'] or 'mijoz')})\n"
            f"📝 {_htm.escape(r['reason'] or 'Pauza')}\n"
            f"🕐 Pauza: {since}" + (f" · {_fmt_until(r['paused_until'])} ga eslatma" if r["paused_until"] else "") + "\n\n"
            f"Davom ettirish uchun mini app'da chatni oching. Keyingi eslatma {_rem_label(r['pause_remind_min'])}dan keyin.")
    await q.add_operator_notification(
        op["id"], "pause", f"Murojaat #{r['id']} pauzada",
        f"{r['full_name'] or 'Mijoz'} · {r['reason'] or 'Pauza'}", r["id"],
        f"pause:{r['id']}:{q.now()[:16]}")
    try:
        if b:
            await b.send_message(op["telegram_id"], text)
    except Exception:
        pass
    await q.set_pause_next_remind(r["id"], r["pause_remind_min"])


async def _pause_expiry_loop():
    """Har daqiqa pauza eslatmalarini yuboradi.

    Pauza muddati avtomatik yakunlanmaydi: operator uni qo'lda davom ettiradi
    yoki mijoz yangi xabar yozganda mavjud oqim pauzani yopadi. Shunda belgilangan
    bir soat o'tgani uchun chat ro'yxatlardan yo'qolib qolmaydi.
    """
    while True:
        await asyncio.sleep(60)
        try:
            for r in await q.due_pause_reminders():
                await _send_pause_reminder(r)
        except Exception:
            logger.exception("pause loop")


# ================================================================
#                 OPERATOR KUNLIK MAQSADI + ADMIN KUNLIK HISOBOTI
# ================================================================
async def operator_goal(op):
    try:
        goal = int(await q.get_setting("op_daily_goal", "0") or 0)
    except ValueError:
        goal = 0
    done = await q.today_done_by_operator(op["id"], now_local().strftime("%Y-%m-%d 00:00:00"))
    return {"goal": goal, "done": done}


def _dur(m):
    m = int(round(m or 0))
    return f"{m} daq" if m < 60 else f"{m // 60} soat {m % 60} daq"


async def daily_report_text(day=None):
    day = day or now_local().date()
    start = day.strftime("%Y-%m-%d 00:00:00")
    end = (day + timedelta(days=1)).strftime("%Y-%m-%d 00:00:00")
    s = await q.day_summary(start, end)
    rows = await q.sla_rows(start, end)
    W = _w()
    first, by_op = [], {}
    for r in rows:
        if r["first_reply"]:
            m = W._mins_between(r["created_at"], r["first_reply"])
            if m is not None and m >= 0:
                first.append(m)
        if r["operator_id"] and r["status"] == "done":
            by_op[r["operator"] or "—"] = by_op.get(r["operator"] or "—", 0) + 1
    top = sorted(by_op.items(), key=lambda x: -x[1])[:3]
    ps = await q.pause_stats(start, end)
    tags = (await q.tag_counts(start, end))[:3]
    lines = [f"📊 <b>Kunlik hisobot — {day.strftime('%d.%m.%Y')}</b>", "",
             f"📥 Murojaatlar: <b>{s['total'] or 0}</b>",
             f"✅ Yakunlangan: <b>{s['done'] or 0}</b>   ❌ Bekor: {s['canceled'] or 0}",
             f"⏳ Hali ochiq: <b>{s['open_'] or 0}</b>" + (f" (shundan ⏸ pauzada: {ps['now']})" if ps["now"] else ""),
             f"⚡ Birinchi javob (median): <b>{_dur(q._median(first)) if first else '—'}</b>"]
    if s["avg_r"]:
        lines.append(f"⭐ O'rtacha baho: <b>{round(s['avg_r'], 1)}</b>" + (f"   ⚠️ past baholar: {s['low']}" if s["low"] else ""))
    if top:
        lines += ["", "🏆 <b>Eng faol operatorlar:</b>"] + [f"{i}. {_htm.escape(n)} — {c} ta" for i, (n, c) in enumerate(top, 1)]
    if tags:
        lines += ["", "🏷 Mavzular: " + ", ".join(f"#{_htm.escape(t)} ({c})" for t, c in tags)]
    return "\n".join(lines)


async def _daily_report_loop():
    while True:
        await asyncio.sleep(30)
        try:
            at = (await q.get_setting("daily_report", "21:00") or "").strip()
            if not at:
                continue
            n = now_local()
            if n.strftime("%H:%M") < at[:5]:
                continue
            today = n.strftime("%Y-%m-%d")
            if await q.get_setting("daily_report_sent", "") == today:
                continue
            await q.set_setting("daily_report_sent", today)
            from utils import cbot
            client = cbot()
            if not client:
                continue
            text = await daily_report_text(n.date())
            for aid in await q.notify_recipient_ids():
                try:
                    await client.send_message(aid, text)
                except Exception:
                    pass
        except Exception:
            logger.exception("daily report")


async def api_admin_daily_preview(request):
    if not await _w()._auth_admin(request, request.query):
        return _json({"ok": False}, 401)
    return _json({"ok": True, "text": await daily_report_text()})


# ================================================================
def register(app: web.Application):
    r = app.router
    r.add_get("/api/products", api_products)
    r.add_post("/api/stock_wait", api_stock_wait)
    r.add_post("/api/pause", api_pause)
    r.add_get("/api/admin/products", api_admin_products)
    r.add_post("/api/admin/product_save", api_admin_product_save)
    r.add_post("/api/admin/product_stock", api_admin_product_stock)
    r.add_post("/api/admin/product_del", api_admin_product_del)
    r.add_post("/api/admin/products_import", api_admin_products_import)
    r.add_get("/api/admin/products_template", api_admin_products_template)
    r.add_post("/api/admin/products_file", api_admin_products_file)
    r.add_get("/api/admin/ai_settings", api_admin_ai_settings)
    r.add_post("/api/admin/ai_settings_save", api_admin_ai_settings_save)
    r.add_post("/api/admin/ai_test", api_admin_ai_test)
    r.add_post("/api/admin/products_ai", api_admin_products_ai)
    r.add_post("/api/admin/products_ai_apply", api_admin_products_ai_apply)
    r.add_get("/api/admin/daily_preview", api_admin_daily_preview)


def start_background():
    asyncio.create_task(_pause_expiry_loop())
    asyncio.create_task(_daily_report_loop())
